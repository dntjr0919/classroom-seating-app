"""
교실 자리 배치 프로그램 - 파일 입출력 및 문서 내보내기 서비스 (services/file_service.py)
- 엑셀(.xlsx, .xls) 학생 명단 불러오기 (헤더 자동 매핑 + 인덱스 폴백)
- 엑셀(.xlsx) 교실 시각 레이아웃 시트 + 학생별 배정 요약 시트 내보내기
- 한글(.hwp) Windows COM 자동화 내보내기 및 비-COM 환경용 PDF / HTML 인쇄 문서 폴백 지원
"""

from __future__ import annotations

import html
import os
import platform
from dataclasses import dataclass
from datetime import datetime
from typing import Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

from models import (
    CellState,
    ClassroomGrid,
    Gender,
    OrderDirection,
    PlacementMode,
    PresetType,
    Seat,
    Student,
    ViewPerspective,
)


class FileServiceError(Exception):
    """파일 입출력 및 내보내기 처리 중 발생하는 예외"""


@dataclass
class RestoredSeatingData:
    """
    엑셀 파일에서 복원된 교실 자리 배치 전체 데이터.
    - grid: 행/열 크기, 좌석별 상태(ACTIVE/AISLE/DISABLED), 고정(is_locked) 및 학생 배정이 복원된 객체
    - roster: 등록된 전체 학생 목록
    - mode: 저장 당시 사용된 자리 배치 방식
    - respect_locks: 저장 당시의 고정석 유지 옵션 상태
    - order_direction: 번호순 배치 시 진행 방향
    - version: 파일 형식 버전
    """
    grid: ClassroomGrid
    roster: List[Student]
    mode: PlacementMode = PlacementMode.RANDOM
    respect_locks: bool = True
    order_direction: OrderDirection = OrderDirection.ROW_FIRST
    version: str = "1.0"


class FileService:
    """
    엑셀 명단 가져오기/내보내기 및 HWP/PDF/HTML 문서 내보내기를 담당하는 서비스 클래스.
    """

    # 엑셀 헤더 자동 매핑 후보 키워드
    NUMBER_HEADERS = {"번호", "학번", "no", "no.", "num", "number", "연번", "출석번호"}
    NAME_HEADERS = {"이름", "성명", "학생명", "학생이름", "name", "student", "성함"}
    GENDER_HEADERS = {"성별", "남녀", "gender", "sex"}

    @classmethod
    def import_students_from_excel(cls, file_path: str) -> List[Student]:
        """
        엑셀 파일(.xlsx, .xls)에서 학생 명단을 읽어 Student 리스트로 반환합니다.
        1차: '번호', '이름', '성별' 헤더 이름을 자동 탐색하여 매핑
        2차(폴백): 일치하는 헤더가 없으면 0번 열=번호, 1번 열=이름, 2번 열=성별로 매핑
        """
        if not os.path.exists(file_path):
            raise FileServiceError(f"파일을 찾을 수 없습니다: {file_path}")

        try:
            df = pd.read_excel(file_path, dtype=str)
        except Exception as exc:
            raise FileServiceError(f"엑셀 파일을 읽는 중 오류가 발생했습니다:\n{exc}") from exc

        if df.empty and len(df.columns) == 0:
            raise FileServiceError("엑셀 파일에 데이터가 없습니다.")

        df = df.fillna("")
        columns = [str(col).strip() for col in df.columns]

        num_col_idx: Optional[int] = None
        name_col_idx: Optional[int] = None
        gender_col_idx: Optional[int] = None

        for idx, col_name in enumerate(columns):
            lowered = col_name.lower().replace(" ", "")
            if num_col_idx is None and lowered in cls.NUMBER_HEADERS:
                num_col_idx = idx
            elif name_col_idx is None and lowered in cls.NAME_HEADERS:
                name_col_idx = idx
            elif gender_col_idx is None and lowered in cls.GENDER_HEADERS:
                gender_col_idx = idx

        # 헤더 매핑 실패 시 열 인덱스 기반 폴백 적용
        used_fallback = False
        if num_col_idx is None and name_col_idx is None:
            used_fallback = True
            if len(columns) >= 2:
                num_col_idx = 0
                name_col_idx = 1
                if len(columns) >= 3:
                    gender_col_idx = 2
            elif len(columns) == 1:
                name_col_idx = 0

        students: List[Student] = []

        # 폴백 모드일 경우 첫 행(columns에 들어간 값)도 유효한 학생 데이터인지 확인
        if used_fallback:
            first_num = columns[num_col_idx] if num_col_idx is not None and not columns[num_col_idx].startswith("Unnamed:") else ""
            first_name = columns[name_col_idx] if name_col_idx is not None and not columns[name_col_idx].startswith("Unnamed:") else ""
            first_gen = columns[gender_col_idx] if gender_col_idx is not None and not columns[gender_col_idx].startswith("Unnamed:") else ""
            if (first_num or first_name) and not (
                first_num.lower() in cls.NUMBER_HEADERS or first_name.lower() in cls.NAME_HEADERS
            ):
                try:
                    students.append(
                        Student(
                            number=first_num,
                            name=first_name,
                            gender=Gender.from_str(first_gen),
                        )
                    )
                except ValueError:
                    pass

        for _, row in df.iterrows():
            vals = [str(v).strip() for v in row.values]
            raw_num = vals[num_col_idx] if (num_col_idx is not None and num_col_idx < len(vals)) else ""
            raw_name = vals[name_col_idx] if (name_col_idx is not None and name_col_idx < len(vals)) else ""
            raw_gender = vals[gender_col_idx] if (gender_col_idx is not None and gender_col_idx < len(vals)) else ""

            # 소수점 문자열 정리 (예: '1.0' -> '1')
            if raw_num.endswith(".0") and raw_num[:-2].isdigit():
                raw_num = raw_num[:-2]

            if not raw_num and not raw_name:
                continue

            try:
                students.append(
                    Student(
                        number=raw_num,
                        name=raw_name,
                        gender=Gender.from_str(raw_gender),
                    )
                )
            except ValueError:
                continue

        if not students:
            raise FileServiceError(
                "유효한 학생 데이터를 찾지 못했습니다. '번호' 또는 '이름' 열이 포함되어 있는지 확인해 주세요."
            )

        return students

    @classmethod
    def export_students_template(cls, file_path: str, students: List[Student]) -> str:
        """
        학생 명단(번호, 이름, 성별)을 깔끔한 엑셀 양식(.xlsx)으로 저장합니다.
        불러오기(import_students_from_excel) 시 100% 호환되는 표준 형식입니다.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "학생명단"
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        thin_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        center_align = Alignment(horizontal="center", vertical="center")

        headers = ["번호", "이름", "성별"]
        ws.row_dimensions[1].height = 28

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = cell_border

        data_font = Font(name="맑은 고딕", size=10)
        for row_idx, st in enumerate(students, start=2):
            ws.row_dimensions[row_idx].height = 22

            gender_str = st.gender.display_kr() if st.gender != Gender.NONE else ""
            c_num = ws.cell(row=row_idx, column=1, value=st.number)
            c_name = ws.cell(row=row_idx, column=2, value=st.name)
            c_gender = ws.cell(row=row_idx, column=3, value=gender_str)

            for cell in (c_num, c_name, c_gender):
                cell.font = data_font
                cell.alignment = center_align
                cell.border = cell_border

        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 12

        try:
            wb.save(file_path)
            return os.path.abspath(file_path)
        except Exception as exc:
            raise FileServiceError(f"엑셀 파일 저장에 실패했습니다:\n{exc}")

    @classmethod
    def _build_layout_sheet(
        cls,
        ws_layout,
        grid: ClassroomGrid,
        title: str = "교실 자리 배치표",
        perspective: ViewPerspective = ViewPerspective.STUDENT_VIEW,
    ) -> None:
        """
        '자리배치표' 시각 레이아웃 시트를 렌더링합니다.
        - STUDENT_VIEW (게시용): 교탁이 위쪽에 위치하고 학생 시점에서 직관적으로 열람
        - TEACHER_VIEW (교사용): 교탁이 아래쪽에 위치하며 180도 회전(행/열 반전, 글자는 정방향)으로 교사 시점에서 열람
        """
        ws_layout.title = "자리배치표"
        ws_layout.views.sheetView[0].showGridLines = True

        is_teacher = (perspective == ViewPerspective.TEACHER_VIEW)
        display_title = f"{title} ({perspective.display_kr()})"

        title_font = Font(name="맑은 고딕", size=16, bold=True, color="1E293B")
        podium_font = Font(name="맑은 고딕", size=12, bold=True, color="FFFFFF")
        desk_font = Font(name="맑은 고딕", size=11, bold=True, color="0F172A")
        empty_desk_font = Font(name="맑은 고딕", size=10, italic=True, color="94A3B8")
        disabled_font = Font(name="맑은 고딕", size=9, color="CBD5E1")

        podium_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        desk_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        locked_desk_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        empty_desk_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        disabled_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        thick_side = Side(border_style="medium", color="1E293B")
        thin_side = Side(border_style="thin", color="CBD5E1")
        desk_border = Border(left=thick_side, right=thick_side, top=thick_side, bottom=thick_side)
        disabled_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        total_cols = max(grid.cols, 1)
        start_col = 2  # B열부터 시작
        end_col = start_col + total_cols - 1

        # 1. 제목 행 (Row 2)
        ws_layout.merge_cells(
            start_row=2, start_column=start_col, end_row=2, end_column=end_col
        )
        title_cell = ws_layout.cell(row=2, column=start_col, value=display_title)
        title_cell.font = title_font
        title_cell.alignment = center_align
        ws_layout.row_dimensions[2].height = 32

        podium_span = min(4, total_cols)
        podium_start = start_col + (total_cols - podium_span) // 2
        podium_end = podium_start + podium_span - 1

        if not is_teacher:
            # 게시용 (교탁 위쪽: Row 4)
            podium_row = 4
            grid_start_row = 6
        else:
            # 교사용 (교탁 아래쪽: 좌석 배치표 아래)
            grid_start_row = 4
            podium_row = grid_start_row + grid.rows + 1

        # 교탁 렌더링 헬퍼
        ws_layout.merge_cells(
            start_row=podium_row, start_column=podium_start, end_row=podium_row, end_column=podium_end
        )
        for c in range(podium_start, podium_end + 1):
            p_cell = ws_layout.cell(row=podium_row, column=c)
            p_cell.fill = podium_fill
            p_cell.border = desk_border
        podium_label = "칠 판 / 교 탁 (앞)" if not is_teacher else "칠 판 / 교 탁 (교사용 - 앞)"
        p_main = ws_layout.cell(row=podium_row, column=podium_start, value=podium_label)
        p_main.font = podium_font
        p_main.alignment = center_align
        ws_layout.row_dimensions[podium_row].height = 28

        # 2. 좌석 격자 렌더링
        for dr in range(grid.rows):
            excel_row = grid_start_row + dr
            ws_layout.row_dimensions[excel_row].height = 42
            for dc in range(grid.cols):
                excel_col = start_col + dc
                # 교사용은 180도 회전 (행/열 반전), 게시용은 정방향
                actual_r = (grid.rows - 1 - dr) if is_teacher else dr
                actual_c = (grid.cols - 1 - dc) if is_teacher else dc

                seat = grid.get_seat(actual_r, actual_c)
                cell = ws_layout.cell(row=excel_row, column=excel_col)

                if seat.state == CellState.ACTIVE:
                    cell.border = desk_border
                    cell.alignment = center_align
                    if seat.student is not None:
                        lock_prefix = "📌 " if seat.is_locked else ""
                        cell.value = f"{lock_prefix}{seat.student.display_label()}"
                        cell.font = desk_font
                        cell.fill = locked_desk_fill if seat.is_locked else desk_fill
                    else:
                        cell.value = "(빈 좌석)"
                        cell.font = empty_desk_font
                        cell.fill = empty_desk_fill
                elif seat.state == CellState.AISLE:
                    cell.value = ""
                else:
                    cell.value = "X"
                    cell.font = disabled_font
                    cell.fill = disabled_fill
                    cell.border = disabled_border
                    cell.alignment = center_align

        ws_layout.column_dimensions["A"].width = 4
        for dc in range(grid.cols):
            col_letter = get_column_letter(start_col + dc)
            actual_c = (grid.cols - 1 - dc) if is_teacher else dc
            is_entire_col_aisle = all(
                grid.get_seat(r, actual_c).state == CellState.AISLE for r in range(grid.rows)
            )
            ws_layout.column_dimensions[col_letter].width = 6 if is_entire_col_aisle else 18

            if is_entire_col_aisle and grid.rows > 1:
                ws_layout.merge_cells(
                    start_row=grid_start_row,
                    start_column=start_col + dc,
                    end_row=grid_start_row + grid.rows - 1,
                    end_column=start_col + dc,
                )

        ws_layout.page_setup.orientation = ws_layout.ORIENTATION_LANDSCAPE
        ws_layout.page_setup.paperSize = ws_layout.PAPERSIZE_A4
        ws_layout.sheet_properties.pageSetUpPr.fitToPage = True

    @classmethod
    def _build_summary_sheet(
        cls,
        ws_summary,
        grid: ClassroomGrid,
        roster: List[Student],
    ) -> None:
        """'학생별 배정명단' 요약 시트를 렌더링합니다."""
        ws_summary.title = "학생별 배정명단"
        headers = ["번호", "이름", "성별", "배정 좌석(행/열)", "고정(Pin) 여부"]
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        body_font = Font(name="맑은 고딕", size=10, color="0F172A")
        thin_side = Side(border_style="thin", color="CBD5E1")
        summary_border = Border(
            left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
        )
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, text in enumerate(headers, start=1):
            h_cell = ws_summary.cell(row=1, column=col_idx, value=text)
            h_cell.fill = header_fill
            h_cell.font = header_font
            h_cell.alignment = center_align
            h_cell.border = summary_border
        ws_summary.row_dimensions[1].height = 26

        assigned_map = grid.assigned_students_map()
        sorted_students = sorted(roster, key=lambda s: s.sort_key())

        for row_idx, student in enumerate(sorted_students, start=2):
            assigned_seat = assigned_map.get(student.student_id)
            coord_str = assigned_seat.coordinate_label() if assigned_seat else "미배정"
            lock_str = "고정됨" if (assigned_seat and assigned_seat.is_locked) else "-"

            row_values = [
                student.number,
                student.name,
                student.gender.display_kr(),
                coord_str,
                lock_str,
            ]
            ws_summary.row_dimensions[row_idx].height = 22
            for col_idx, val in enumerate(row_values, start=1):
                c_cell = ws_summary.cell(row=row_idx, column=col_idx, value=val)
                c_cell.font = body_font
                c_cell.alignment = center_align
                c_cell.border = summary_border

        col_widths = [10, 16, 10, 18, 14]
        for idx, width in enumerate(col_widths, start=1):
            ws_summary.column_dimensions[get_column_letter(idx)].width = width

    @classmethod
    def _build_metadata_sheet(
        cls,
        ws_meta,
        grid: ClassroomGrid,
        roster: List[Student],
        mode: PlacementMode,
        respect_locks: bool,
        order_direction: OrderDirection,
    ) -> None:
        """
        '배치복원_상세데이터' 시트를 렌더링합니다.
        프로그램에서 오차 없이 정확한 왕복 복원(Round-trip restoration)을 수행할 수 있도록
        메타데이터, 격자 셀 상태, 학생 배정 정보를 규격화된 테이블로 기록합니다.
        """
        ws_meta.title = "배치복원_상세데이터"

        header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
        section_font = Font(name="맑은 고딕", size=11, bold=True, color="1E293B")
        data_font = Font(name="맑은 고딕", size=10, color="0F172A")
        header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        thin_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        # 1. 메타데이터 섹션
        meta_items = [
            ("FORMAT_VERSION", "1.0", "파일 형식 버전", "교실 자리 배치 앱 규격 버전"),
            ("ROWS", str(grid.rows), "격자 행 수", "교실 격자 N 행"),
            ("COLS", str(grid.cols), "격자 열 수", "교실 격자 M 열"),
            ("PLACEMENT_MODE", mode.value, "배치 방식", mode.display_kr()),
            ("RESPECT_LOCKS", "TRUE" if respect_locks else "FALSE", "고정석 유지 옵션", "배치 시 고정석 유지 여부"),
            ("ORDER_DIRECTION", order_direction.value, "번호순 배치 방향", order_direction.display_kr()),
            ("SAVED_AT", datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "저장 일시", "파일 생성 시각"),
        ]

        ws_meta.cell(row=1, column=1, value="[시스템 메타데이터 및 배치 설정]").font = section_font
        meta_headers = ["설정 KEY", "설정 VALUE", "항목명", "설명"]
        for c_idx, h_text in enumerate(meta_headers, start=1):
            c = ws_meta.cell(row=2, column=c_idx, value=h_text)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center_align
            c.border = cell_border

        for r_offset, (k, v, name, desc) in enumerate(meta_items, start=3):
            ws_meta.cell(row=r_offset, column=1, value=k).alignment = left_align
            ws_meta.cell(row=r_offset, column=2, value=v).alignment = center_align
            ws_meta.cell(row=r_offset, column=3, value=name).alignment = left_align
            ws_meta.cell(row=r_offset, column=4, value=desc).alignment = left_align
            for c_idx in range(1, 5):
                cell = ws_meta.cell(row=r_offset, column=c_idx)
                cell.font = data_font
                cell.border = cell_border

        # 2. 격자 셀 상태 섹션 (ACTIVE / AISLE / DISABLED 및 고정 여부)
        grid_start_r = len(meta_items) + 4
        ws_meta.cell(row=grid_start_r, column=1, value="[격자 셀 상태]").font = section_font
        grid_headers = ["행(0-based)", "열(0-based)", "좌표라벨", "셀상태(ACTIVE/AISLE/DISABLED)", "고정여부(TRUE/FALSE)"]
        for c_idx, h_text in enumerate(grid_headers, start=1):
            c = ws_meta.cell(row=grid_start_r + 1, column=c_idx, value=h_text)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center_align
            c.border = cell_border

        curr_r = grid_start_r + 2
        for r in range(grid.rows):
            for c in range(grid.cols):
                seat = grid.get_seat(r, c)
                row_vals = [
                    r,
                    c,
                    seat.coordinate_label(),
                    seat.state.value,
                    "TRUE" if seat.is_locked else "FALSE",
                ]
                for c_idx, val in enumerate(row_vals, start=1):
                    cell = ws_meta.cell(row=curr_r, column=c_idx, value=val)
                    cell.font = data_font
                    cell.alignment = center_align
                    cell.border = cell_border
                curr_r += 1

        # 3. 학생 명단 및 좌석 배정 섹션
        student_start_r = curr_r + 1
        ws_meta.cell(row=student_start_r, column=1, value="[학생 명단 및 좌석 배정]").font = section_font
        st_headers = [
            "학생ID", "번호", "이름", "성별(M/F/NONE)", "배정행(0-based)", "배정열(0-based)", "배정좌표라벨", "고정여부(TRUE/FALSE)"
        ]
        for c_idx, h_text in enumerate(st_headers, start=1):
            c = ws_meta.cell(row=student_start_r + 1, column=c_idx, value=h_text)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center_align
            c.border = cell_border

        assigned_map = grid.assigned_students_map()
        sorted_students = sorted(roster, key=lambda s: s.sort_key())

        curr_r = student_start_r + 2
        for student in sorted_students:
            seat = assigned_map.get(student.student_id)
            if seat is not None:
                row_vals = [
                    student.student_id,
                    student.number,
                    student.name,
                    student.gender.value,
                    seat.row,
                    seat.col,
                    seat.coordinate_label(),
                    "TRUE" if seat.is_locked else "FALSE",
                ]
            else:
                row_vals = [
                    student.student_id,
                    student.number,
                    student.name,
                    student.gender.value,
                    "",
                    "",
                    "미배정",
                    "FALSE",
                ]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_meta.cell(row=curr_r, column=c_idx, value=val)
                cell.font = data_font
                cell.alignment = center_align
                cell.border = cell_border
            curr_r += 1

        # 열 너비 정리
        ws_meta.column_dimensions["A"].width = 22
        ws_meta.column_dimensions["B"].width = 16
        ws_meta.column_dimensions["C"].width = 16
        ws_meta.column_dimensions["D"].width = 30
        ws_meta.column_dimensions["E"].width = 16
        ws_meta.column_dimensions["F"].width = 16
        ws_meta.column_dimensions["G"].width = 16
        ws_meta.column_dimensions["H"].width = 18

    @classmethod
    def export_to_excel(
        cls,
        file_path: str,
        grid: ClassroomGrid,
        roster: List[Student],
        title: str = "교실 자리 배치표",
        perspective: ViewPerspective = ViewPerspective.STUDENT_VIEW,
    ) -> str:
        """
        현재 교실 자리 배치를 엑셀(.xlsx) 파일로 내보냅니다 (게시 및 인쇄 열람용).
        - 시트 1 ('자리배치표'): 교탁 및 교실 시각적 레이아웃 (게시용: 교탁 위 / 교사용: 교탁 아래 180도 회전)
        - 시트 2 ('학생별 배정명단'): 번호/이름순 학생 목록과 배정 좌석 좌표 요약표
        """
        wb = openpyxl.Workbook()
        ws_layout = wb.active
        cls._build_layout_sheet(ws_layout, grid, title, perspective)
        ws_summary = wb.create_sheet()
        cls._build_summary_sheet(ws_summary, grid, roster)

        try:
            wb.save(file_path)
        except Exception as exc:
            raise FileServiceError(f"엑셀 파일 저장 중 오류가 발생했습니다:\n{exc}") from exc

        return file_path

    @classmethod
    def save_seating_to_excel(
        cls,
        file_path: str,
        grid: ClassroomGrid,
        roster: List[Student],
        mode: PlacementMode = PlacementMode.RANDOM,
        respect_locks: bool = True,
        order_direction: OrderDirection = OrderDirection.ROW_FIRST,
        title: str = "교실 자리 배치표",
    ) -> str:
        """
        현재 교실 자리 배치와 학생 명단, 레이아웃 및 설정 전체를 복원 가능한 엑셀(.xlsx) 파일로 저장합니다.
        시트 1 ('자리배치표'): 사람이 읽기 좋은 시각적 배치표
        시트 2 ('학생별 배정명단'): 학생별 배정 요약표
        시트 3 ('배치복원_상세데이터'): 프로그램 복원에 필요한 정밀 메타데이터/격자셀/학생배정 데이터
        """
        wb = openpyxl.Workbook()
        ws_layout = wb.active
        cls._build_layout_sheet(ws_layout, grid, title)
        ws_summary = wb.create_sheet()
        cls._build_summary_sheet(ws_summary, grid, roster)
        ws_meta = wb.create_sheet()
        cls._build_metadata_sheet(ws_meta, grid, roster, mode, respect_locks, order_direction)

        try:
            wb.save(file_path)
        except Exception as exc:
            raise FileServiceError(f"자리 배치 엑셀 파일 저장 중 오류가 발생했습니다:\n{exc}") from exc

        return file_path

    @classmethod
    def load_seating_from_excel(cls, file_path: str) -> RestoredSeatingData:
        """
        본 앱에서 저장한 자리 배치 엑셀 파일(.xlsx)을 읽어와 무결성을 검증한 후
        ClassroomGrid 및 Student 명단, 배치 상태, 설정을 완전히 복원하여 반환합니다.
        오류 발견 시 기존 애플리케이션 상태를 건드리지 않고 FileServiceError를 발생시킵니다 (Atomic).
        """
        if not os.path.exists(file_path):
            raise FileServiceError(f"파일을 찾을 수 없습니다: {file_path}")

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as exc:
            raise FileServiceError(f"엑셀 파일을 열 수 없습니다:\n{exc}") from exc

        meta_sheet_name = "배치복원_상세데이터"
        if meta_sheet_name not in wb.sheetnames:
            raise FileServiceError(
                f"선택한 엑셀 파일에 복원용 데이터 시트('{meta_sheet_name}')가 없습니다.\n\n"
                "이 기능은 본 앱에서 [자리 배치 엑셀 저장]으로 생성한 파일만 복원할 수 있습니다.\n"
                "단순 학생 명단만 가져오려면 상단의 [학생 명단 관리]를 이용해 주세요."
            )

        ws = wb[meta_sheet_name]

        # 1. 메타데이터 파싱
        meta: Dict[str, str] = {}
        for r in range(2, 12):
            k = str(ws.cell(row=r, column=1).value or "").strip().upper()
            v = str(ws.cell(row=r, column=2).value or "").strip()
            if k:
                meta[k] = v

        version = meta.get("FORMAT_VERSION", "")
        if version != "1.0":
            raise FileServiceError(
                f"지원하지 않거나 알 수 없는 자리 배치 파일 형식 버전입니다: '{version}' (지원 버전: 1.0)"
            )

        try:
            rows = int(meta.get("ROWS", "0"))
            cols = int(meta.get("COLS", "0"))
        except ValueError:
            raise FileServiceError("격자 크기(행/열) 정보가 올바른 정수 형식이 아닙니다.")

        if rows < 1 or cols < 1 or rows > 50 or cols > 50:
            raise FileServiceError(f"지원 범위를 벗어난 격자 크기입니다: {rows}행 × {cols}열 (최대 50×50)")

        mode_str = meta.get("PLACEMENT_MODE", PlacementMode.RANDOM.value)
        try:
            mode = PlacementMode(mode_str)
        except ValueError:
            mode = PlacementMode.RANDOM

        respect_locks = meta.get("RESPECT_LOCKS", "TRUE").upper() in ("TRUE", "1", "Y", "YES", "참")

        order_dir_str = meta.get("ORDER_DIRECTION", OrderDirection.ROW_FIRST.value)
        try:
            order_direction = OrderDirection(order_dir_str)
        except ValueError:
            order_direction = OrderDirection.ROW_FIRST

        # 2. 새로운 ClassroomGrid 인스턴스 준비
        grid = ClassroomGrid(rows=rows, cols=cols)

        # 3. 셀 상태 파싱
        # ws 내에서 '[격자 셀 상태]' 및 '[학생 명단 및 좌석 배정]' 섹션 위치 탐색
        grid_section_row = -1
        student_section_row = -1

        for r in range(1, ws.max_row + 1):
            val = str(ws.cell(row=r, column=1).value or "").strip()
            if val.startswith("[격자 셀 상태]"):
                grid_section_row = r
            elif val.startswith("[학생 명단 및 좌석 배정]"):
                student_section_row = r
                break

        if grid_section_row == -1 or student_section_row == -1:
            raise FileServiceError("엑셀 파일 내 필수 데이터 섹션([격자 셀 상태] 또는 [학생 명단])을 찾을 수 없습니다.")

        # 셀 상태 읽기
        for r in range(grid_section_row + 2, student_section_row):
            raw_r = ws.cell(row=r, column=1).value
            raw_c = ws.cell(row=r, column=2).value
            if raw_r is None or raw_c is None:
                continue

            try:
                cell_r = int(raw_r)
                cell_c = int(raw_c)
            except ValueError:
                continue

            if not (0 <= cell_r < rows and 0 <= cell_c < cols):
                raise FileServiceError(f"좌석 좌표가 격자 범위를 초과합니다: ({cell_r + 1}행, {cell_c + 1}열)")

            state_str = str(ws.cell(row=r, column=4).value or "ACTIVE").strip()
            lock_str = str(ws.cell(row=r, column=5).value or "").strip().upper()

            try:
                cell_state = CellState(state_str)
            except ValueError:
                cell_state = CellState.ACTIVE

            seat = grid.get_seat(cell_r, cell_c)
            seat.set_state(cell_state)
            if cell_state == CellState.ACTIVE:
                seat.is_locked = lock_str in ("TRUE", "1", "Y", "YES", "고정", "고정됨")

        # 4. 학생 명단 및 좌석 배정 파싱
        roster: List[Student] = []
        seen_student_ids: Set[str] = set()
        assigned_seat_coords: Set[Tuple[int, int]] = set()

        for r in range(student_section_row + 2, ws.max_row + 1):
            raw_num = str(ws.cell(row=r, column=2).value or "").strip()
            raw_name = str(ws.cell(row=r, column=3).value or "").strip()
            raw_gen = str(ws.cell(row=r, column=4).value or "").strip()
            raw_ar = ws.cell(row=r, column=5).value
            raw_ac = ws.cell(row=r, column=6).value
            raw_locked = str(ws.cell(row=r, column=8).value or "").strip().upper()

            # 번호 소수점 제거 (예: '1.0' -> '1')
            if raw_num.endswith(".0") and raw_num[:-2].isdigit():
                raw_num = raw_num[:-2]

            if not raw_num and not raw_name:
                continue

            try:
                student = Student(
                    number=raw_num,
                    name=raw_name,
                    gender=Gender.from_str(raw_gen),
                )
            except ValueError as ve:
                raise FileServiceError(f"학생 데이터 행({r}행)에 유효하지 않은 정보가 있습니다: {ve}")

            if student.student_id in seen_student_ids:
                raise FileServiceError(
                    f"엑셀 파일 내에 중복된 학생 식별자가 존재합니다: {student.display_label()}\n"
                    "동명이인의 경우 출석번호를 다르게 지정해야 식별할 수 있습니다."
                )
            seen_student_ids.add(student.student_id)
            roster.append(student)

            # 좌석 배정 정보가 있는 경우 검증 및 배정
            if raw_ar is not None and raw_ac is not None and str(raw_ar).strip() != "" and str(raw_ac).strip() != "":
                try:
                    ar = int(raw_ar)
                    ac = int(raw_ac)
                except ValueError:
                    raise FileServiceError(
                        f"학생 {student.display_label()}의 배정 좌표({raw_ar}, {raw_ac})가 정수가 아닙니다."
                    )

                if not (0 <= ar < rows and 0 <= ac < cols):
                    raise FileServiceError(
                        f"학생 {student.display_label()}의 배정 좌석({ar + 1}행 {ac + 1}열)이 "
                        f"교실 격자 크기({rows}행 × {cols}열)를 벗어납니다."
                    )

                target_seat = grid.get_seat(ar, ac)
                if not target_seat.is_active_desk:
                    raise FileServiceError(
                        f"학생 {student.display_label()}이 통로 또는 비활성 영역({target_seat.coordinate_label()})에 "
                        "배정될 수 없습니다."
                    )

                if (ar, ac) in assigned_seat_coords:
                    prev_name = target_seat.student.display_label() if target_seat.student else "기존 학생"
                    raise FileServiceError(
                        f"좌석 {target_seat.coordinate_label()}에 둘 이상의 학생({prev_name}, {student.display_label()})이 "
                        "중복 배정되어 있습니다."
                    )

                assigned_seat_coords.add((ar, ac))
                target_seat.student = student
                if raw_locked in ("TRUE", "1", "Y", "YES", "고정", "고정됨"):
                    target_seat.is_locked = True

        return RestoredSeatingData(
            grid=grid,
            roster=roster,
            mode=mode,
            respect_locks=respect_locks,
            order_direction=order_direction,
            version=version,
        )

    @classmethod
    def create_example_seating_template(
        cls,
        file_path: str,
        rows: int = 5,
        cols: int = 8,
    ) -> str:
        """외부에서 직접 편집 가능한 예제 자리 배치 엑셀 서식을 생성합니다."""
        sample_grid = ClassroomGrid(rows=rows, cols=cols)
        sample_grid.apply_preset(PresetType.PAIRS)
        sample_roster = [
            Student("1", "강하늘", Gender.MALE),
            Student("2", "고윤정", Gender.FEMALE),
            Student("3", "김도현", Gender.MALE),
            Student("4", "김민지", Gender.FEMALE),
        ]
        sample_grid.get_seat(0, 0).student = sample_roster[0]
        sample_grid.get_seat(0, 0).is_locked = True
        sample_grid.get_seat(0, 1).student = sample_roster[1]

        return cls.save_seating_to_excel(
            file_path,
            sample_grid,
            sample_roster,
            mode=PlacementMode.MIXED_GENDER_PAIR,
            respect_locks=True,
            order_direction=OrderDirection.ROW_FIRST,
            title="교실 자리 배치표 (예제 서식)",
        )


    @classmethod
    def export_document_with_fallback(
        cls,
        file_path: str,
        grid: ClassroomGrid,
        roster: List[Student],
        title: str = "교실 자리 배치표",
        perspective: ViewPerspective = ViewPerspective.STUDENT_VIEW,
    ) -> Tuple[str, str]:
        """
        한글(.hwp) 또는 인쇄용 문서를 내보냅니다.
        1순위: 사용자가 .hwp 확장자를 선택하고 Windows + pywin32 + 한컴오피스 한글이 설치된 환경이면
              COM 자동화를 통해 실제 .hwp 표 문서를 생성합니다.
        2순위(폴백): 비-Windows 환경이거나 한컴오피스/pywin32가 없는 경우 프로그램이 중단되지 않도록
              PDF(.pdf) 또는 인쇄 최적화 HTML(.html) 문서로 안전하게 내보냅니다.

        반환값: (실제 저장된 파일 경로, 내보내기 결과 안내 메시지)
        """
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".hwp":
            # Windows + pywin32 + HWP COM 시도
            if platform.system() == "Windows":
                try:
                    saved_path = cls._export_via_hwp_com(file_path, grid, title, perspective)
                    return saved_path, f"한컴오피스 한글(.hwp) 문서({perspective.display_kr()})가 성공적으로 생성되었습니다."
                except Exception as hwp_exc:
                    # 폴백: 동일 경로에 .pdf 또는 .html 생성
                    fallback_pdf = os.path.splitext(file_path)[0] + ".pdf"
                    try:
                        saved_pdf = cls.export_to_pdf(fallback_pdf, grid, roster, title, perspective)
                        return (
                            saved_pdf,
                            f"현재 PC에서 한컴오피스 한글(HWP) 자동화를 실행할 수 없어({hwp_exc}), "
                            f"즉시 인쇄 가능한 PDF 문서({perspective.display_kr()})로 대체 저장했습니다:\n{saved_pdf}",
                        )
                    except Exception:
                        fallback_html = os.path.splitext(file_path)[0] + ".html"
                        saved_html = cls.export_to_html(fallback_html, grid, roster, title, perspective)
                        return (
                            saved_html,
                            f"한글(HWP) 프로그램이 감지되지 않아 인쇄용 HTML 문서({perspective.display_kr()})로 대체 저장했습니다:\n{saved_html}",
                        )
            else:
                # 비-Windows 환경에서 .hwp 요청 시 PDF 또는 HTML로 안전 폴백
                fallback_pdf = os.path.splitext(file_path)[0] + ".pdf"
                try:
                    saved_pdf = cls.export_to_pdf(fallback_pdf, grid, roster, title, perspective)
                    return (
                        saved_pdf,
                        f"현재 운영체제({platform.system()})에서는 한컴오피스 COM을 사용할 수 없어 "
                        f"인쇄용 PDF 문서({perspective.display_kr()})로 자동 변환하여 저장했습니다:\n{saved_pdf}",
                    )
                except Exception:
                    fallback_html = os.path.splitext(file_path)[0] + ".html"
                    saved_html = cls.export_to_html(fallback_html, grid, roster, title, perspective)
                    return (
                        saved_html,
                        f"현재 환경에서 HWP COM을 사용할 수 없어 인쇄용 HTML 문서({perspective.display_kr()})로 저장했습니다:\n{saved_html}",
                    )

        if ext == ".pdf":
            try:
                saved_pdf = cls.export_to_pdf(file_path, grid, roster, title, perspective)
                return saved_pdf, f"인쇄용 PDF 문서({perspective.display_kr()})가 성공적으로 저장되었습니다."
            except Exception:
                fallback_html = os.path.splitext(file_path)[0] + ".html"
                saved_html = cls.export_to_html(fallback_html, grid, roster, title, perspective)
                return (
                    saved_html,
                    f"PDF 폰트/렌더러 제약으로 인해 브라우저 인쇄 최적화 HTML 파일({perspective.display_kr()})로 저장했습니다:\n{saved_html}",
                )

        # 기본 HTML 내보내기
        if ext not in (".html", ".htm"):
            file_path = file_path + ".html"
        saved_html = cls.export_to_html(file_path, grid, roster, title, perspective)
        return saved_html, f"브라우저 및 한글/워드에서 열 수 있는 인쇄용 HTML 문서({perspective.display_kr()})가 저장되었습니다."

    @classmethod
    def _export_via_hwp_com(
        cls,
        file_path: str,
        grid: ClassroomGrid,
        title: str = "교실 자리 배치표",
        perspective: ViewPerspective = ViewPerspective.STUDENT_VIEW,
    ) -> str:
        """
        Windows 환경에서 win32com.client를 사용하여 한컴오피스 한글(.hwp) 문서를 생성합니다.
        교사용 선택 시 180도 회전(행/열 반전, 글자는 정방향) 및 교탁 아래 배치.
        """
        import win32com.client  # type: ignore

        hwp = win32com.client.Dispatch("HWPFrame.HwpObject")
        try:
            hwp.RegisterModule("FilePathCheckDLL", "FilePathCheckerModule")
        except Exception:
            pass

        hwp.HAction.Run("FileNew")

        is_teacher = (perspective == ViewPerspective.TEACHER_VIEW)
        display_title = f"{title} ({perspective.display_kr()})"

        # 문서 제목 입력
        act = hwp.CreateAction("InsertText")
        pset = act.CreateSet()
        if not is_teacher:
            pset.SetItem("Text", f"{display_title}\r\n[칠판 / 교탁 (앞)]\r\n\r\n")
        else:
            pset.SetItem("Text", f"{display_title}\r\n\r\n")
        act.Execute(pset)

        # 표 생성 (grid.rows x grid.cols)
        tp = hwp.HParameterSet.HTableCreation
        hwp.HAction.GetDefault("TableCreate", tp.HSet)
        tp.Rows = grid.rows
        tp.Cols = grid.cols
        hwp.HAction.Execute("TableCreate", tp.HSet)

        for dr in range(grid.rows):
            for dc in range(grid.cols):
                actual_r = (grid.rows - 1 - dr) if is_teacher else dr
                actual_c = (grid.cols - 1 - dc) if is_teacher else dc
                seat = grid.get_seat(actual_r, actual_c)
                text = ""
                if seat.state == CellState.ACTIVE:
                    text = seat.student.display_label() if seat.student else "(빈 좌석)"
                elif seat.state == CellState.AISLE:
                    text = " "
                else:
                    text = "-"

                ins = hwp.CreateAction("InsertText")
                ins_set = ins.CreateSet()
                ins_set.SetItem("Text", text)
                ins.Execute(ins_set)

                if not (dr == grid.rows - 1 and dc == grid.cols - 1):
                    hwp.HAction.Run("TableRightCell")

        if is_teacher:
            hwp.HAction.Run("MoveDown")
            act_bottom = hwp.CreateAction("InsertText")
            pset_bottom = act_bottom.CreateSet()
            pset_bottom.SetItem("Text", "\r\n[칠판 / 교탁 (교사용 - 앞)]\r\n")
            act_bottom.Execute(pset_bottom)

        abs_path = os.path.abspath(file_path)
        hwp.SaveAs(abs_path, "HWP")
        hwp.Quit()
        return abs_path

    @classmethod
    def export_to_pdf(
        cls,
        file_path: str,
        grid: ClassroomGrid,
        roster: List[Student],
        title: str = "교실 자리 배치표",
        perspective: ViewPerspective = ViewPerspective.STUDENT_VIEW,
    ) -> str:
        """
        ReportLab을 사용하여 A4 가로 방향 PDF 자리 배치표를 생성합니다.
        교사용 선택 시 180도 회전(행/열 반전, 글자는 정방향) 및 교탁 아래 배치.
        """
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        font_name = cls._register_korean_font_for_reportlab(pdfmetrics, TTFont)
        is_teacher = (perspective == ViewPerspective.TEACHER_VIEW)
        display_title = f"{title} ({perspective.display_kr()})"

        doc = SimpleDocTemplate(
            file_path,
            pagesize=landscape(A4),
            leftMargin=28,
            rightMargin=28,
            topMargin=28,
            bottomMargin=28,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "KoreanTitle",
            parent=styles["Title"],
            fontName=font_name,
            fontSize=18,
            leading=24,
            alignment=1,
            textColor=colors.HexColor("#1E293B"),
        )
        sub_style = ParagraphStyle(
            "KoreanSub",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=11,
            leading=16,
            alignment=1,
            textColor=colors.HexColor("#FFFFFF"),
            backColor=colors.HexColor("#334155"),
        )

        podium_text = "칠 판 / 교 탁 (앞)" if not is_teacher else "칠 판 / 교 탁 (교사용 - 앞)"
        podium_widget = Paragraph(podium_text, sub_style)

        elements = [
            Paragraph(html.escape(display_title), title_style),
            Spacer(1, 10),
        ]
        if not is_teacher:
            elements.extend([podium_widget, Spacer(1, 14)])

        table_data: List[List[str]] = []
        table_cmds = [
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]

        for dr in range(grid.rows):
            row_cells: List[str] = []
            for dc in range(grid.cols):
                actual_r = (grid.rows - 1 - dr) if is_teacher else dr
                actual_c = (grid.cols - 1 - dc) if is_teacher else dc
                seat = grid.get_seat(actual_r, actual_c)

                if seat.state == CellState.ACTIVE:
                    if seat.student:
                        prefix = "[고정] " if seat.is_locked else ""
                        row_cells.append(f"{prefix}{seat.student.display_label()}")
                    else:
                        row_cells.append("(빈 좌석)")
                    bg = colors.HexColor("#FEF3C7") if seat.is_locked else colors.HexColor("#F8FAFC")
                    table_cmds.append(("BACKGROUND", (dc, dr), (dc, dr), bg))
                    table_cmds.append(("BOX", (dc, dr), (dc, dr), 1.5, colors.HexColor("#1E293B")))
                elif seat.state == CellState.AISLE:
                    row_cells.append("")
                else:
                    row_cells.append("X")
                    table_cmds.append(("BACKGROUND", (dc, dr), (dc, dr), colors.HexColor("#F1F5F9")))
                    table_cmds.append(("TEXTCOLOR", (dc, dr), (dc, dr), colors.HexColor("#94A3B8")))
            table_data.append(row_cells)

        avail_width = landscape(A4)[0] - 56
        col_w = avail_width / max(grid.cols, 1)
        row_h = min(46, max(28, (landscape(A4)[1] - 150) / max(grid.rows, 1)))

        t = Table(table_data, colWidths=[col_w] * grid.cols, rowHeights=[row_h] * grid.rows)
        t.setStyle(TableStyle(table_cmds))
        elements.append(t)

        if is_teacher:
            elements.extend([Spacer(1, 14), podium_widget])

        doc.build(elements)
        return file_path

    @staticmethod
    def _register_korean_font_for_reportlab(pdfmetrics, TTFont) -> str:
        """시스템에 설치된 한글 TTF/TTC 폰트를 탐색하여 ReportLab에 등록합니다."""
        candidate_paths = [
            "C:/Windows/Fonts/malgun.ttf",
            "C:/Windows/Fonts/gulim.ttc",
            "C:/Windows/Fonts/batang.ttc",
            "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/noto/NotoSansKR-Regular.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        ]
        for path in candidate_paths:
            if os.path.exists(path):
                try:
                    if path.lower().endswith(".ttc"):
                        pdfmetrics.registerFont(TTFont("KoreanFont", path, subfontIndex=0))
                    else:
                        pdfmetrics.registerFont(TTFont("KoreanFont", path))
                    return "KoreanFont"
                except Exception:
                    continue
        return "Helvetica"

    @classmethod
    def export_to_html(
        cls,
        file_path: str,
        grid: ClassroomGrid,
        roster: List[Student],
        title: str = "교실 자리 배치표",
        perspective: ViewPerspective = ViewPerspective.STUDENT_VIEW,
    ) -> str:
        """
        모든 환경에서 깨짐 없이 바로 인쇄하거나 한글(HWP)/MS Word에서 불러올 수 있는
        고해상도 A4 가로 인쇄 최적화 HTML 문서를 생성합니다.
        교사용 선택 시 180도 회전(행/열 반전, 글자는 정방향) 및 교탁 아래 배치.
        """
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
        is_teacher = (perspective == ViewPerspective.TEACHER_VIEW)
        display_title = f"{title} ({perspective.display_kr()})"

        rows_html: List[str] = []

        for dr in range(grid.rows):
            cells_html: List[str] = []
            for dc in range(grid.cols):
                actual_r = (grid.rows - 1 - dr) if is_teacher else dr
                actual_c = (grid.cols - 1 - dc) if is_teacher else dc
                seat = grid.get_seat(actual_r, actual_c)

                if seat.state == CellState.ACTIVE:
                    if seat.student:
                        lock_badge = "<span class='pin'>📌</span> " if seat.is_locked else ""
                        cls_name = "desk locked" if seat.is_locked else "desk"
                        label = html.escape(seat.student.display_label())
                        coord = html.escape(seat.coordinate_label())
                        cells_html.append(
                            f"<td class='{cls_name}'><div class='coord'>{coord}</div>"
                            f"<div class='student'>{lock_badge}{label}</div></td>"
                        )
                    else:
                        coord = html.escape(seat.coordinate_label())
                        cells_html.append(
                            f"<td class='desk empty'><div class='coord'>{coord}</div>"
                            f"<div class='student-empty'>(빈 좌석)</div></td>"
                        )
                elif seat.state == CellState.AISLE:
                    cells_html.append("<td class='aisle'></td>")
                else:
                    cells_html.append("<td class='disabled'>비활성</td>")
            rows_html.append("<tr>" + "".join(cells_html) + "</tr>")

        grid_table_body = "\n".join(rows_html)

        podium_label = "칠 판 / 교 탁 (앞)" if not is_teacher else "칠 판 / 교 탁 (교사용 - 앞)"
        podium_div = f"<div class='podium'>{podium_label}</div>"

        top_podium = podium_div if not is_teacher else ""
        bottom_podium = podium_div if is_teacher else ""

        html_doc = f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="UTF-8">
<title>{html.escape(display_title)}</title>
<style>
  @page {{ size: A4 landscape; margin: 12mm; }}
  body {{
    font-family: "Malgun Gothic", "맑은 고딕", "Apple SD Gothic Neo", "Noto Sans KR", sans-serif;
    color: #0F172A;
    margin: 0;
    padding: 20px;
    background: #FFFFFF;
  }}
  .header {{ text-align: center; margin-bottom: 16px; }}
  .header h1 {{ margin: 0 0 6px 0; font-size: 24px; color: #1E293B; }}
  .header .meta {{ font-size: 12px; color: #64748B; }}
  .podium {{
    width: 280px;
    margin: 10px auto;
    padding: 10px 16px;
    background: #334155;
    color: #FFFFFF;
    font-weight: 700;
    text-align: center;
    border-radius: 6px;
    letter-spacing: 2px;
  }}
  table.classroom {{
    width: 100%;
    border-collapse: separate;
    border-spacing: 8px;
    table-layout: fixed;
    margin: 12px 0;
  }}
  td.desk {{
    border: 2px solid #1E293B;
    border-radius: 6px;
    background: #F8FAFC;
    height: 64px;
    text-align: center;
    vertical-align: middle;
    padding: 6px;
  }}
  td.desk.locked {{
    background: #FEF3C7;
    border-color: #D97706;
  }}
  td.desk.empty {{
    background: #FFFFFF;
    border-style: dashed;
    border-color: #94A3B8;
  }}
  td.aisle {{
    border: none;
    background: transparent;
  }}
  td.disabled {{
    border: 1px solid #E2E8F0;
    background: #F1F5F9;
    color: #94A3B8;
    font-size: 11px;
    text-align: center;
  }}
  .coord {{ font-size: 10px; color: #64748B; margin-bottom: 4px; }}
  .student {{ font-size: 14px; font-weight: 700; color: #0F172A; word-break: keep-all; }}
  .student-empty {{ font-size: 12px; color: #94A3B8; }}
</style>
</head>
<body>
  <div class="header">
    <h1>{html.escape(display_title)}</h1>
    <div class="meta">생성 일시: {now_str} | 출력 방향: {perspective.display_kr()} | 격자 크기: {grid.rows}행 × {grid.cols}열 | 활성 좌석: {len(grid.active_desks())}석 / 등록 학생: {len(roster)}명</div>
  </div>
  {top_podium}
  <table class="classroom">
    {grid_table_body}
  </table>
  {bottom_podium}
</body>
</html>
"""
        try:
            with open(file_path, "w", encoding="utf-8") as f:
                f.write(html_doc)
        except Exception as exc:
            raise FileServiceError(f"HTML 문서 저장 중 오류가 발생했습니다:\n{exc}") from exc

        return file_path

